from __future__ import annotations

import argparse
import hashlib
import re
from collections import defaultdict
from collections.abc import Iterable
from pathlib import Path
from typing import Any

import pandas as pd
import yaml
from openpyxl import load_workbook
from pyxlsb import open_workbook

FROZEN_YEARS = ["2018/19", "2019/20", "2020/21", "2021/22", "2022/23", "2023/24", "2024/25"]
FIELDS = {"05", "06", "07"}
MIN_YEARS = 5
INITIAL_CYCLES = {
    "Curso técnico superior profissional",
    "Licenciatura 1.º ciclo",
    "Mestrado integrado",
}
FIRST_CYCLE = {"Licenciatura 1.º ciclo"}
SECOND_CYCLE = {"Mestrado 2.º ciclo", "Mestrado integrado", "Mestrado integrado terminal"}
DOCTORAL_ENROLMENT = {"Doutoramento 3.º ciclo"}
DOCTORAL_GRADUATE = {"Doutoramento 3.º ciclo", "Doutoramento"}
COMPONENTS = [
    "first_time_entrants",
    "graduates_first_cycle",
    "graduates_second_cycle",
    "doctoral_enrolments",
    "doctoral_graduates",
]


def _normalise_code(value: Any, width: int) -> str | None:
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return None
    text = str(value).strip()
    if not text:
        return None
    try:
        text = str(int(float(text)))
    except ValueError:
        pass
    return text.zfill(width)


def _short_year(value: str) -> str:
    match = re.search(r"(20\d{2})[/_](20\d{2})", value)
    if not match:
        raise ValueError(f"Cannot identify academic year from {value!r}")
    return f"{match.group(1)}/{match.group(2)[-2:]}"


def _sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as fh:
        for chunk in iter(lambda: fh.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def _xlsb_rows(path: Path) -> Iterable[dict[str, Any]]:
    with open_workbook(path) as book, book.get_sheet("Tabela1") as sheet:
        header: list[Any] | None = None
        index: dict[str, int] = {}
        for row in sheet.rows():
            values = [cell.v for cell in row]
            if header is None:
                if "Código do estabelecimento de ensino" in values:
                    header = values
                    index = {str(v).strip(): i for i, v in enumerate(values) if v is not None}
                continue

            def get(name: str) -> Any:
                i = index[name]
                return values[i] if i < len(values) else None

            yield {
                "academic_year": _short_year(str(get("Ano letivo"))),
                "institution_code": _normalise_code(
                    get("Código do estabelecimento de ensino"), 4
                ),
                "institution_name": get("Estabelecimento de ensino"),
                "cycle": get("Curso/Ciclo de estudos"),
                "field": _normalise_code(
                    get("Área de educação e formação - Código da área geral"), 2
                ),
                "first_time": get("Primeira vez"),
                "count": get("N"),
            }


def _choose_diplomados_sheet(book: Any) -> str:
    candidates = [
        sheet
        for sheet in book.sheetnames
        if sheet.startswith("agregados_") or sheet.startswith("Diplomados20")
    ]
    if len(candidates) != 1:
        raise ValueError(f"Expected one Diplomados data sheet, found {candidates}")
    return candidates[0]


def _xlsx_rows(path: Path) -> Iterable[dict[str, Any]]:
    book = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet_name = _choose_diplomados_sheet(book)
        ws = book[sheet_name]
        rows = ws.iter_rows(values_only=True)
        header = [str(v).strip() if v is not None else "" for v in next(rows)]
        index = {name: i for i, name in enumerate(header) if name}
        year = _short_year(sheet_name)
        count_col = "Total" if "Total" in index else "N"
        for values in rows:
            code = _normalise_code(values[index["Código Estabelecimento"]], 4)
            if code is None:
                continue
            yield {
                "academic_year": year,
                "institution_code": code,
                "institution_name": values[index["Estabelecimento"]],
                "cycle": values[index["Curso/Ciclo de Estudos"]],
                "field": _normalise_code(
                    values[index["Área de Educação e Formação - Código Área Geral"]], 2
                ),
                "count": values[index[count_col]],
            }
    finally:
        book.close()


def _aggregate_inscritos(input_dir: Path) -> tuple[pd.DataFrame, dict[str, str]]:
    records: dict[tuple[str, str, str, str], float] = defaultdict(float)
    names: dict[tuple[str, str], str] = {}
    hashes: dict[str, str] = {}
    found: set[str] = set()
    for path in sorted(input_dir.glob("*.xlsb")):
        year_seen: str | None = None
        for row in _xlsb_rows(path):
            year = row["academic_year"]
            year_seen = year
            if (
                year not in FROZEN_YEARS
                or row["field"] not in FIELDS
                or row["institution_code"] is None
            ):
                continue
            found.add(year)
            code = row["institution_code"]
            names[(year, code)] = str(row["institution_name"])
            cycle = str(row["cycle"])
            n = float(row["count"] or 0)
            if row["first_time"] == "S" and cycle in INITIAL_CYCLES:
                records[(year, code, row["field"], "first_time_entrants")] += n
            if cycle in DOCTORAL_ENROLMENT:
                records[(year, code, row["field"], "doctoral_enrolments")] += n
        if year_seen in FROZEN_YEARS:
            hashes[year_seen] = _sha256(path)
    if found != set(FROZEN_YEARS):
        raise ValueError(f"Inscritos years mismatch: found {sorted(found)}")
    out = pd.DataFrame(
        [
            {
                "academic_year": year,
                "institution_code": code,
                "institution_name": names.get((year, code), ""),
                "field": field,
                "component": component,
                "count": count,
            }
            for (year, code, field, component), count in records.items()
        ]
    )
    return out, hashes


def _aggregate_diplomados(input_dir: Path) -> tuple[pd.DataFrame, dict[str, str]]:
    records: dict[tuple[str, str, str, str], float] = defaultdict(float)
    names: dict[tuple[str, str], str] = {}
    hashes: dict[str, str] = {}
    found: set[str] = set()
    for path in sorted(input_dir.glob("*.xlsx")):
        year_for_file: str | None = None
        for row in _xlsx_rows(path):
            year = row["academic_year"]
            year_for_file = year
            if year not in FROZEN_YEARS or row["field"] not in FIELDS:
                continue
            found.add(year)
            code = row["institution_code"]
            names[(year, code)] = str(row["institution_name"])
            cycle = str(row["cycle"])
            n = float(row["count"] or 0)
            component = None
            if cycle in FIRST_CYCLE:
                component = "graduates_first_cycle"
            elif cycle in SECOND_CYCLE:
                component = "graduates_second_cycle"
            elif cycle in DOCTORAL_GRADUATE:
                component = "doctoral_graduates"
            if component:
                records[(year, code, row["field"], component)] += n
        if year_for_file in FROZEN_YEARS:
            hashes[year_for_file] = _sha256(path)
    if found != set(FROZEN_YEARS):
        raise ValueError(f"Diplomados years mismatch: found {sorted(found)}")
    out = pd.DataFrame(
        [
            {
                "academic_year": year,
                "institution_code": code,
                "institution_name": names.get((year, code), ""),
                "field": field,
                "component": component,
                "count": count,
            }
            for (year, code, field, component), count in records.items()
        ]
    )
    return out, hashes


def _support_outputs(
    inscritos: pd.DataFrame, diplomados: pd.DataFrame
) -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    counts = pd.concat([inscritos, diplomados], ignore_index=True)
    support = (
        counts.groupby(["institution_code", "field", "component"], as_index=False)
        .agg(supported_years=("academic_year", "nunique"))
    )
    latest_names = (
        counts.sort_values("academic_year")
        .drop_duplicates("institution_code", keep="last")
        .set_index("institution_code")["institution_name"]
    )
    support.insert(1, "institution_name", support["institution_code"].map(latest_names))
    support["meets_five_year_minimum"] = support["supported_years"] >= MIN_YEARS
    support = support.sort_values(["institution_code", "field", "component"]).reset_index(
        drop=True
    )

    pair = support.pivot_table(
        index=["institution_code", "institution_name", "field"],
        columns="component",
        values="supported_years",
        aggfunc="first",
    ).reset_index()
    for component in COMPONENTS:
        if component not in pair:
            pair[component] = pd.NA
    pair["eligible_components"] = sum(
        pair[component].fillna(0).ge(MIN_YEARS).astype(int) for component in COMPONENTS
    )
    pair["all_five_eligible"] = pair["eligible_components"].eq(5)
    pair = pair[
        [
            "institution_code",
            "institution_name",
            "field",
            *COMPONENTS,
            "eligible_components",
            "all_five_eligible",
        ]
    ].sort_values(["institution_code", "field"])

    ins_codes = set(inscritos["institution_code"])
    dip_codes = set(diplomados["institution_code"])
    concordance_rows = []
    for code in sorted(ins_codes | dip_codes):
        institution_rows = counts[counts["institution_code"].eq(code)].sort_values("academic_year")
        canonical_name = str(institution_rows.iloc[-1]["institution_name"])
        aliases = sorted(set(institution_rows["institution_name"].astype(str)) - {canonical_name})
        concordance_rows.append(
            {
                "institution_code": code,
                "canonical_name": canonical_name,
                "present_inscritos": code in ins_codes,
                "present_diplomados": code in dip_codes,
                "aliases": " | ".join(aliases),
                "join_rule": "numeric_establishment_code",
            }
        )
    concordance = pd.DataFrame(concordance_rows)
    return support, pair.reset_index(drop=True), concordance


def _audit(
    support: pd.DataFrame, pair: pd.DataFrame, concordance: pd.DataFrame
) -> dict[str, Any]:
    component_summary = []
    for component, group in support.groupby("component", sort=True):
        component_summary.append(
            {
                "component": component,
                "series_total": len(group),
                "series_meeting_five_year_minimum": int(
                    group["meets_five_year_minimum"].sum()
                ),
                "median_supported_years": float(group["supported_years"].median()),
                "supported_year_distribution": {
                    int(k): int(v)
                    for k, v in group["supported_years"].value_counts().sort_index().items()
                },
            }
        )
    return {
        "version": 1,
        "study": "C",
        "audit_type": "raides_institution_field_component_support",
        "status": "raides_support_gate_passed_fct_institution_weights_pending",
        "support_rule": {
            "minimum_comparable_annual_observations": MIN_YEARS,
            "year_is_supported_only_when_actual_record_exists": True,
            "absent_row_is_not_promoted_to_zero": True,
            "fields": sorted(FIELDS),
        },
        "summary": {
            "institution_field_component_series": len(support),
            "eligible_series": int(support["meets_five_year_minimum"].sum()),
            "institution_field_pairs": len(pair),
            "pairs_with_all_five_components_eligible": int(pair["all_five_eligible"].sum()),
            "institution_codes": len(concordance),
            "codes_in_both_raides_families": int(
                (concordance["present_inscritos"] & concordance["present_diplomados"]).sum()
            ),
        },
        "component_summary": component_summary,
        "eligible_component_distribution": {
            int(k): int(v)
            for k, v in pair["eligible_components"].value_counts().sort_index().items()
        },
        "field_summary": [
            {
                "field": str(field),
                "institution_field_pairs": len(group),
                "pairs_with_all_five_components_eligible": int(
                    group["all_five_eligible"].sum()
                ),
                "median_eligible_components": float(group["eligible_components"].median()),
            }
            for field, group in pair.groupby("field")
        ],
        "institution_code_concordance": {
            "inscritos_codes": int(concordance["present_inscritos"].sum()),
            "diplomados_codes": int(concordance["present_diplomados"].sum()),
            "shared_codes": int(
                (concordance["present_inscritos"] & concordance["present_diplomados"]).sum()
            ),
            "inscritos_only_codes": concordance.loc[
                concordance["present_inscritos"] & ~concordance["present_diplomados"],
                "institution_code",
            ].tolist(),
            "diplomados_only_codes": concordance.loc[
                ~concordance["present_inscritos"] & concordance["present_diplomados"],
                "institution_code",
            ].tolist(),
            "names_are_labels_not_join_keys": True,
        },
        "unit_exposure_allowed": False,
        "remaining_primary_blockers": [
            "formal_fct_unit_institution_participation_not_resolved",
            "unit_institution_field_weights_not_built",
        ],
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--inscritos-dir", type=Path, required=True)
    parser.add_argument("--diplomados-dir", type=Path, required=True)
    parser.add_argument("--output-dir", type=Path, required=True)
    args = parser.parse_args()

    inscritos, _ = _aggregate_inscritos(args.inscritos_dir)
    diplomados, _ = _aggregate_diplomados(args.diplomados_dir)
    support, pair, concordance = _support_outputs(inscritos, diplomados)
    args.output_dir.mkdir(parents=True, exist_ok=True)
    support.to_csv(args.output_dir / "study_c_support_matrix.csv", index=False)
    pair.to_csv(args.output_dir / "study_c_support_by_institution_field.csv", index=False)
    concordance.to_csv(args.output_dir / "study_c_institution_concordance.csv", index=False)
    (args.output_dir / "raides_support_audit.yml").write_text(
        yaml.safe_dump(_audit(support, pair, concordance), sort_keys=False, allow_unicode=True),
        encoding="utf-8",
    )


if __name__ == "__main__":
    main()
